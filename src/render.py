"""Stage 3: tailored JSON -> ATS-safe .docx, plus the HTML digest.

ATS rules obeyed here: single column, no tables, no text boxes, no headers or
footers, no images, standard section names, standard font. Parsers choke on
everything else.
"""
from __future__ import annotations

import html
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "work"
OUT = ROOT / "out"

DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55)


def _style(doc):
    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(10.5)
    n.font.color.rgb = DARK
    pf = n.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.06
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.55)
        s.left_margin = s.right_margin = Inches(0.65)


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(10)
    pb = p.paragraph_format
    pb.keep_with_next = True
    # thin rule under the heading, via paragraph border
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    pPr.append(borders)
    return p


def _bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    return p


def build_resume(data: dict, basics: dict, path: Path):
    doc = Document()
    _style(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(basics.get("name", ""))
    r.bold = True
    r.font.size = Pt(17)

    contact = " | ".join(
        x for x in [
            basics.get("location"), basics.get("email"), basics.get("phone"),
            basics.get("linkedin"), basics.get("github"),
        ] if x
    )
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = c.add_run(contact)
    cr.font.size = Pt(9)
    cr.font.color.rgb = GREY

    if data.get("summary"):
        _heading(doc, "Summary")
        doc.add_paragraph(data["summary"])

    if data.get("skills"):
        _heading(doc, "Skills")
        for label, items in data["skills"]:
            sp = doc.add_paragraph()
            b = sp.add_run(f"{label}: ")
            b.bold = True
            sp.add_run(items)

    if data.get("experience"):
        _heading(doc, "Experience")
        for job in data["experience"]:
            jp = doc.add_paragraph()
            jp.paragraph_format.space_before = Pt(5)
            jp.paragraph_format.keep_with_next = True
            jr = jp.add_run(f"{job.get('role','')} — {job.get('company','')}")
            jr.bold = True
            meta = " | ".join(x for x in [job.get("location"), job.get("dates")] if x)
            if meta:
                mp = doc.add_paragraph()
                mp.paragraph_format.keep_with_next = True
                mr = mp.add_run(meta)
                mr.italic = True
                mr.font.size = Pt(9)
                mr.font.color.rgb = GREY
            for b in job.get("bullets", []):
                _bullet(doc, b)

    if data.get("projects"):
        _heading(doc, "Projects")
        for pr in data["projects"]:
            pp = doc.add_paragraph()
            pp.paragraph_format.keep_with_next = True
            n = pp.add_run(pr.get("name", ""))
            n.bold = True
            if pr.get("link"):
                lk = pp.add_run(f"  {pr['link']}")
                lk.font.size = Pt(9)
                lk.font.color.rgb = GREY
            for b in pr.get("bullets", []):
                _bullet(doc, b)

    if data.get("education"):
        _heading(doc, "Education")
        for ed in data["education"]:
            ep = doc.add_paragraph()
            er = ep.add_run(f"{ed.get('credential','')} — {ed.get('institution','')}")
            er.bold = True
            tail = " | ".join(x for x in [ed.get("year"), ed.get("note")] if x)
            if tail:
                t = ep.add_run(f"  {tail}")
                t.font.size = Pt(9)
                t.font.color.rgb = GREY

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


CSS = """
body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;
     background:#f4f4f2;margin:0;padding:24px;color:#1a1a1a}
.wrap{max-width:680px;margin:0 auto}
h1{font-size:19px;margin:0 0 4px}
.sub{color:#666;font-size:13px;margin:0 0 20px}
.card{background:#fff;border:1px solid #e2e2de;border-radius:10px;
      padding:16px 18px;margin-bottom:12px}
.top{display:flex;justify-content:space-between;align-items:baseline;gap:12px}
.title{font-weight:600;font-size:15px;margin:0}
.co{color:#444;font-size:13px;margin:2px 0 8px}
.score{font-weight:700;font-size:15px;white-space:nowrap}
.s-hi{color:#1a7f37}.s-mid{color:#9a6700}.s-lo{color:#888}
.tag{display:inline-block;font-size:11px;padding:2px 8px;border-radius:20px;
     background:#eef1ee;color:#444;margin:0 5px 5px 0}
.tag.flag{background:#fdf0d5;color:#8a6100}
.why{font-size:13px;color:#333;margin:6px 0}
.gaps{font-size:12px;color:#8a6100;margin:6px 0 0}
a.btn{display:inline-block;margin-top:10px;margin-right:8px;font-size:13px;
      text-decoration:none;color:#0969da;font-weight:500}
.foot{color:#888;font-size:12px;margin-top:22px;text-align:center}
"""


def build_digest(rows: list[dict], resume_files: dict, stats: dict) -> str:
    today = datetime.now().strftime("%A, %d %B %Y")
    parts = [f"<html><head><meta charset='utf-8'><style>{CSS}</style></head><body><div class='wrap'>"]
    parts.append(f"<h1>AI/ML remote roles — {len(rows)} new</h1>")
    parts.append(f"<p class='sub'>{today} &middot; {stats.get('raw',0)} postings scanned "
                 f"&middot; {len(resume_files)} resumes tailored</p>")

    if not rows:
        parts.append("<div class='card'>Nothing new cleared the filters today. "
                     "That's normal on weekends.</div>")

    for r in rows:
        sc = r.get("score", 0)
        cls = "s-hi" if sc >= 75 else ("s-mid" if sc >= 60 else "s-lo")
        parts.append("<div class='card'>")
        parts.append(
            f"<div class='top'><p class='title'>{html.escape(r.get('title',''))}</p>"
            f"<span class='score {cls}'>{sc}</span></div>"
        )
        parts.append(f"<p class='co'>{html.escape(r.get('company',''))} &middot; "
                     f"{html.escape(r.get('location',''))}</p>")
        elig = r.get("eligibility", "unclear")
        flag = " flag" if elig == "unclear" else ""
        parts.append(f"<span class='tag{flag}'>{elig}</span>")
        if r.get("id") in resume_files:
            parts.append("<span class='tag'>resume ready</span>")
        if r.get("why"):
            parts.append(f"<p class='why'>{html.escape(r['why'])}</p>")
        if r.get("gaps"):
            parts.append("<p class='gaps'>Gaps: " +
                         html.escape(", ".join(r["gaps"])) + "</p>")
        parts.append(f"<a class='btn' href='{r.get('url','#')}'>Open posting &rarr;</a>")
        parts.append("</div>")

    parts.append("<p class='foot'>Tailored resumes are attached as .docx. "
                 "Check every claim before you send it.</p>")
    parts.append("</div></body></html>")
    return "".join(parts)


def main():
    import yaml
    screened = json.loads((WORK / "screened.json").read_text()) if (WORK / "screened.json").exists() else []
    basics = yaml.safe_load((ROOT / "profile" / "resume.yml").read_text()).get("basics", {})

    resume_files = {}
    tdir = WORK / "tailored"
    if tdir.exists():
        for f in sorted(tdir.glob("*.json")):
            data = json.loads(f.read_text())
            safe_co = "".join(ch for ch in data.get("company", "co") if ch.isalnum() or ch in " -_").strip()
            safe_ti = "".join(ch for ch in data.get("title", "role") if ch.isalnum() or ch in " -_").strip()
            name = basics.get("name", "Resume").replace(" ", "_")
            out = OUT / "resumes" / f"{name}_{safe_co}_{safe_ti}.docx"[:180]
            try:
                build_resume(data, basics, out)
                resume_files[data["id"]] = out
            except Exception as e:
                print(f"  resume render failed for {f.name}: {e}")

    stats = {}
    if (WORK / "new_jobs.json").exists():
        stats["raw"] = len(json.loads((WORK / "new_jobs.json").read_text()))

    OUT.mkdir(exist_ok=True)
    (OUT / "digest.html").write_text(build_digest(screened, resume_files, stats))
    (OUT / "attachments.json").write_text(json.dumps([str(p) for p in resume_files.values()]))
    print(f"Digest: {len(screened)} jobs, {len(resume_files)} resumes")


if __name__ == "__main__":
    main()
